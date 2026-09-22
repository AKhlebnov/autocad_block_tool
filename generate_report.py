# ============================================================
#  generate_report.py – построение отчёта в Excel.
#  Версия 3.0
#  Читает один общий export_final.csv и строит 4 листа:
#    1. Сводка по шкафам
#    2. Детали по панелям
#    3. Порты (с автофильтром)
#    4. Кабельный журнал (см. cable_journal.py)
# ============================================================

import sys               # работа с системой
import os                # пути к файлам
import importlib.util    # динамическая загрузка config.py
import time              # для замера времени
import pandas as pd      # для чтения CSV и записи Excel
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

# Импортируем модуль кабельного журнала
from cable_journal import add_cable_journal_sheet


# ------------------------------------------------------------
# Константы оформления
# ------------------------------------------------------------

# Светло-голубой фон для шапок (приятный «акцентный» цвет Excel)
HEADER_FILL = PatternFill(
    start_color="D9E1F2",
    end_color="D9E1F2",
    fill_type="solid"
)


# ------------------------------------------------------------
# Вспомогательные функции для форматирования Excel
# ------------------------------------------------------------

def apply_borders(ws, row_start, row_end, col_start, col_end):
    """
    Применяет тонкие границы ко всем ячейкам в указанном диапазоне.

    Параметры:
        ws        — рабочий лист openpyxl
        row_start — первая строка (1-индексация)
        row_end   — последняя строка (включительно)
        col_start — первый столбец (1-индексация)
        col_end   — последний столбец (включительно)
    """
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    for row in range(row_start, row_end + 1):
        for col in range(col_start, col_end + 1):
            ws.cell(row=row, column=col).border = thin_border


def style_header(ws, num_cols, header_row=1):
    """
    Оформляет заголовки таблицы:
      - жирный шрифт,
      - выравнивание по центру с переносом,
      - светло-голубой фон.

    Параметры:
        ws          — рабочий лист
        num_cols    — количество столбцов (ширина шапки)
        header_row  — номер строки заголовка (по умолчанию 1)
    """
    for col_idx in range(1, num_cols + 1):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.fill = HEADER_FILL


def auto_column_widths(ws, widths):
    """Устанавливает ширину столбцов по словарю {буква_столбца: ширина}."""
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width


def add_auto_filter(ws, num_cols, last_row):
    """
    Включает автофильтр на указанном листе.

    Параметры:
        ws        — рабочий лист
        num_cols  — количество столбцов
        last_row  — номер последней строки данных
    """
    last_col_letter = get_column_letter(num_cols)
    ws.auto_filter.ref = f"A1:{last_col_letter}{last_row}"


# ------------------------------------------------------------
# Основная функция
# ------------------------------------------------------------

def main(wait_for_exit=True):
    """
    Основная функция отчёта:
    - читает export_final.csv,
    - строит листы «Сводка по шкафам», «Детали по панелям», «Порты»,
    - добавляет лист «Кабельный журнал»,
    - сохраняет результат в «Отчёт по проекту.xlsx».
    """
    start_total = time.time()

    # --- 1. Загрузка config (для порядка типов блоков) ---
    if getattr(sys, 'frozen', False):
        application_path = os.path.dirname(sys.executable)
    else:
        application_path = os.path.dirname(os.path.abspath(__file__))
    config_path = os.path.join(application_path, 'config.py')
    spec = importlib.util.spec_from_file_location("config", config_path)
    config = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(config)

    # Словарь порядка типов блоков: {имя_блока: индекс}
    # Нужен, чтобы сортировать «Тип оборудования» не по алфавиту,
    # а по порядку из настроек (как в таблице GUI).
    order_map = {}
    idx = 0
    for b in getattr(config, "BLOCK_CONFIGS", []):
        bname = b.get("block_name", "")
        if bname and bname not in order_map:
            order_map[bname] = idx
            idx += 1
    # Добавляем блоки особого режима (после основных)
    if getattr(config, "SPECIAL_MODE_ENABLED", False):
        for name in getattr(config, "SPECIAL_BLOCK_NAMES", []):
            if name and name not in order_map:
                order_map[name] = idx
                idx += 1

    def type_order(dev_type):
        """Возвращает порядковый номер типа блока для сортировки."""
        return order_map.get(dev_type, 999)

    # --- 2. Входной и выходной файлы ---
    INPUT_CSV = "export_final.csv"
    OUTPUT_EXCEL = "Отчёт по проекту.xlsx"

    if not os.path.exists(INPUT_CSV):
        print(f"Ошибка: файл {INPUT_CSV} не найден.")
        print("Сначала запустите экспорт и расчёт имён (assign_names.py).")
        if wait_for_exit:
            input("\nНажмите Enter для выхода...")
        sys.exit(1)

    print(f"Читаем файл: {INPUT_CSV}")

    # --- 3. Чтение данных ---
    df = pd.read_csv(INPUT_CSV, encoding="utf-8-sig")

    required_cols = ["Шкаф", "Этаж", "Панель", "Порт", "Имя", "Тип"]
    for col in required_cols:
        if col not in df.columns:
            print(f"Ошибка: в файле {INPUT_CSV} нет колонки '{col}'.")
            if wait_for_exit:
                input("\nНажмите Enter для выхода...")
            sys.exit(1)

    print(f"Записей в файле: {len(df)}")

    # Приводим типы
    df["Панель"] = df["Панель"].astype(int)
    df["Порт"] = df["Порт"].astype(int)
    df["Этаж"] = pd.to_numeric(df["Этаж"], errors="coerce").fillna(0).astype(int)

    # ============================================================
    # 4. Лист «Сводка по шкафам»
    # ============================================================
    summary_data = []

    for cabinet, group in df.groupby("Шкаф"):
        total_devices = len(group)
        panels = sorted(group["Панель"].unique())
        num_panels = len(panels)

        panel_details = []
        for panel in panels:
            panel_group = group[group["Панель"] == panel]
            ports = sorted(panel_group["Порт"].tolist())

            # Сжимаем список портов в диапазоны
            ranges = []
            start = ports[0]
            end = ports[0]
            for p in ports[1:]:
                if p == end + 1:
                    end = p
                else:
                    ranges.append(f"{start}-{end}" if start != end else str(start))
                    start = end = p
            ranges.append(f"{start}-{end}" if start != end else str(start))

            # Типы на панели — сортируем по порядку из настроек
            types = sorted(panel_group["Тип"].unique(), key=type_order)
            types_str = ", ".join(types)

            panel_details.append(
                f"Панель {panel}: {','.join(ranges)} (типы: {types_str})"
            )

        # Каждую панель — на новой строке
        summary_data.append({
            "Шкаф": cabinet,
            "Всего устройств": total_devices,
            "Количество патч-панелей": num_panels,
            "Детали по панелям": "\n".join(panel_details),
        })

    df_summary = pd.DataFrame(summary_data)

    # ============================================================
    # 5. Лист «Детали по панелям»
    # ============================================================
    detail_rows = []
    for cabinet, group in df.groupby("Шкаф"):
        for panel, panel_group in group.groupby("Панель"):
            ports = sorted(panel_group["Порт"].tolist())

            ranges = []
            start = ports[0]
            end = ports[0]
            for p in ports[1:]:
                if p == end + 1:
                    end = p
                else:
                    ranges.append(f"{start}-{end}" if start != end else str(start))
                    start = end = p
            ranges.append(f"{start}-{end}" if start != end else str(start))
            port_str = ",".join(ranges)

            # Типы — по порядку из настроек
            types = sorted(panel_group["Тип"].unique(), key=type_order)
            types_str = ", ".join(types)

            detail_rows.append({
                "Шкаф": cabinet,
                "Номер панели": panel,
                "Занятые порты": port_str,
                "Количество портов": len(ports),
                "Тип оборудования": types_str,
            })

    df_detail = pd.DataFrame(detail_rows)

    # ============================================================
    # 6. Лист «Порты»
    # ============================================================
    port_rows = []
    MAX_PORT = 48

    for (cabinet, panel), panel_group in df.groupby(["Шкаф", "Панель"]):
        taken = {}
        for _, row in panel_group.iterrows():
            taken[int(row["Порт"])] = (
                row["Имя"],
                int(row["Этаж"]),
                row["Тип"],
            )

        for port in range(1, MAX_PORT + 1):
            if port in taken:
                name, floor, dev_type = taken[port]
            else:
                name, floor, dev_type = "Свободный порт", "", ""
            port_rows.append({
                "Шкаф": cabinet,
                "Панель": panel,
                "Порт": port,
                "Этаж": floor,
                "Оборудование": name,
                "Тип оборудования": dev_type,
            })

    df_ports = pd.DataFrame(port_rows)

    # ============================================================
    # 7. Запись в Excel
    # ============================================================
    try:
        with pd.ExcelWriter(OUTPUT_EXCEL, engine="openpyxl") as writer:
            # --- Три стандартных листа ---
            df_summary.to_excel(writer, sheet_name="Сводка по шкафам", index=False)
            df_detail.to_excel(writer, sheet_name="Детали по панелям", index=False)
            df_ports.to_excel(writer, sheet_name="Порты", index=False)

            # ----------------------------------------------------
            # 7.1. Оформление листа «Сводка по шкафам»
            # ----------------------------------------------------
            ws = writer.sheets["Сводка по шкафам"]
            auto_column_widths(ws, {
                'A': 18,   # Шкаф
                'B': 16,   # Всего устройств
                'C': 24,   # Количество патч-панелей
                'D': 75,   # Детали по панелям
            })
            # Шапка со стилем (жирный + цвет + по центру)
            style_header(ws, num_cols=4)
            # Данные: выравнивание
            for row_idx in range(2, len(df_summary) + 2):
                for col_idx in (1, 2, 3):
                    ws.cell(row=row_idx, column=col_idx).alignment = Alignment(
                        horizontal='center', vertical='center', wrap_text=True
                    )
                ws.cell(row=row_idx, column=4).alignment = Alignment(
                    horizontal='left', vertical='top', wrap_text=True
                )
            apply_borders(ws, 1, len(df_summary) + 1, 1, 4)

            # ----------------------------------------------------
            # 7.2. Оформление листа «Детали по панелям»
            # ----------------------------------------------------
            ws = writer.sheets["Детали по панелям"]
            auto_column_widths(ws, {
                'A': 18,   # Шкаф
                'B': 14,   # Номер панели
                'C': 30,   # Занятые порты
                'D': 16,   # Количество портов
                'E': 35,   # Тип оборудования
            })
            style_header(ws, num_cols=5)
            for row_idx in range(2, len(df_detail) + 2):
                for col_idx in (1, 2, 4):
                    ws.cell(row=row_idx, column=col_idx).alignment = Alignment(
                        horizontal='center', vertical='center', wrap_text=True
                    )
                for col_idx in (3, 5):
                    ws.cell(row=row_idx, column=col_idx).alignment = Alignment(
                        horizontal='left', vertical='center', wrap_text=True
                    )
            apply_borders(ws, 1, len(df_detail) + 1, 1, 5)

            # ----------------------------------------------------
            # 7.3. Оформление листа «Порты»
            # ----------------------------------------------------
            ws = writer.sheets["Порты"]
            auto_column_widths(ws, {
                'A': 18,   # Шкаф
                'B': 10,   # Панель
                'C': 10,   # Порт
                'D': 10,   # Этаж
                'E': 26,   # Оборудование
                'F': 22,   # Тип оборудования
            })
            style_header(ws, num_cols=6)
            for row_idx in range(2, len(df_ports) + 2):
                for col_idx in (1, 2, 3, 4):
                    ws.cell(row=row_idx, column=col_idx).alignment = Alignment(
                        horizontal='center', vertical='center'
                    )
                for col_idx in (5, 6):
                    ws.cell(row=row_idx, column=col_idx).alignment = Alignment(
                        horizontal='left', vertical='center'
                    )
            apply_borders(ws, 1, len(df_ports) + 1, 1, 6)
            # === АВТОФИЛЬТР на листе «Порты» ===
            add_auto_filter(ws, num_cols=6, last_row=len(df_ports) + 1)

            # ----------------------------------------------------
            # 7.4. Лист «Кабельный журнал» — строит отдельный модуль
            # ----------------------------------------------------
            df_sorted = df.sort_values(by=["Шкаф", "Панель", "Порт"]).reset_index(drop=True)
            add_cable_journal_sheet(writer, df_sorted)

        print(f"Отчёт сохранён в '{OUTPUT_EXCEL}'")

    except PermissionError:
        print(f"\n❌ Ошибка: не удалось сохранить файл '{OUTPUT_EXCEL}'.")
        print("Возможно, файл уже открыт в Excel. Закройте его и повторите.")
    except Exception as e:
        print(f"\n❌ Ошибка при создании отчёта: {e}")
        import traceback
        traceback.print_exc()

    elapsed = time.time() - start_total
    print(f"\nОбщее время выполнения: {elapsed:.2f} сек")

    if wait_for_exit:
        input("\n✅ Готово! Нажмите Enter для выхода...")


if __name__ == "__main__":
    main(wait_for_exit=True)
