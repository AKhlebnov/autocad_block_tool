# ============================================================
#  cross_table.py – генерация листа «Кроссировочная таблица».
#  Показывает подключения портов патч-панели к портам коммутатора.
#  Версия 4.2
#
#  Структура:
#    Строка 1: общие шапки («Телекоммуникационный шкаф», 
#              «Оконечное оборудование», «Примечание»).
#    Строка 2: конкретные заголовки столбцов.
#    Данные: с 3-й строки, по одной строке на каждый порт 1..48.
# ============================================================

from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


# Светло-голубой фон для шапки — тот же, что и на остальных листах
HEADER_FILL = PatternFill(
    start_color="D9E1F2",
    end_color="D9E1F2",
    fill_type="solid"
)


def switch_port_from_patch(patch_port):
    """
    Преобразует номер порта патч-панели в номер порта коммутатора
    по правилу:
      • порты 1-24 патч-панели  → НЕЧЁТНЫЕ порты коммутатора (1, 3, ..., 47)
      • порты 25-48 патч-панели → ЧЁТНЫЕ порты коммутатора (2, 4, ..., 48)

    Формула:
      1-24:  switch = patch * 2 - 1
      25-48: switch = (patch - 24) * 2

    Примеры:
      1  → 1
      2  → 3
      24 → 47
      25 → 2
      26 → 4
      48 → 48
    """
    if patch_port <= 24:
        return patch_port * 2 - 1
    else:
        return (patch_port - 24) * 2


def add_cross_table_sheet(writer, df, block_configs):
    """
    Добавляет в книгу Excel лист «Кроссировочная таблица».

    Столбцы:
        A: №
        --- Телекоммуникационный шкаф ---
        B: Шкаф №
        C: Коммутатор №
        D: Порт коммутатора №
        E: Патч-панель №
        F: Порт патч-панели №
        --- Оконечное оборудование ---
        G: Тип оборудования
        H: Маркировка по проекту
        I: Этаж расположения
        J: Питание
        K: IP адрес
        L: MAC адрес
        --- отдельно ---
        M: Примечание

    Правила заполнения:
        • Свободные (резервные) порты помечаются как «резерв».
        • Все пустые ячейки заполняются прочерком «-».
        • Столбец «Примечание» остаётся пустым (его заполняет подрядчик).
        • Питание: «PoE», если для типа блока стоит галочка PoE, иначе «нет».
        • IP и MAC — «-» (заполнит подрядчик после монтажа).

    Параметры:
        writer        — pd.ExcelWriter (openpyxl), уже открытый
        df            — DataFrame из export_final.csv с колонками:
                        Шкаф, Этаж, Панель, Порт, Имя, Тип, _display_name
        block_configs — список словарей BLOCK_CONFIGS из config.py
                        (нужен для определения PoE по типу блока)
    """
    wb = writer.book
    ws = wb.create_sheet("Кроссировочная таблица")

    # --- Стили ---
    header_font = Font(bold=True, size=11)
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # --- Словарь PoE: {имя_блока: True/False} ---
    poe_map = {}
    for b in block_configs:
        bname = b.get("block_name", "")
        if bname:
            poe_map[bname] = bool(b.get("poe", False))

    # ============================================================
    # ШАПКА (2 строки)
    # ============================================================
    # --- Строка 1: общие шапки для групп столбцов ---
    ws.cell(row=1, column=1, value="№")
    ws.cell(row=1, column=2, value="Телекоммуникационный шкаф")
    ws.cell(row=1, column=7, value="Оконечное оборудование")
    ws.cell(row=1, column=13, value="Примечание")

    # --- Строка 2: конкретные заголовки столбцов ---
    # Столбец A (№) и M (Примечание) объединяются со строкой 1,
    # поэтому у них второй строки нет.
    headers_row2 = {
        2:  "Шкаф №",
        3:  "Коммутатор №",
        4:  "Порт коммутатора №",
        5:  "Патч-панель №",
        6:  "Порт патч-панели №",
        7:  "Тип оборудования",
        8:  "Маркировка по проекту",
        9:  "Этаж расположения",
        10: "Питание",
        11: "IP адрес",
        12: "MAC адрес",
    }
    for col, header in headers_row2.items():
        ws.cell(row=2, column=col, value=header)

    # --- Объединение ячеек шапки ---
    ws.merge_cells('A1:A2')       # № — на две строки
    ws.merge_cells('B1:F1')       # «Телекоммуникационный шкаф» — B:F в строке 1
    ws.merge_cells('G1:L1')       # «Оконечное оборудование» — G:L в строке 1
    ws.merge_cells('M1:M2')       # «Примечание» — на две строки

    # --- Стилизация шапки (строки 1 и 2) ---
    for row in (1, 2):
        for col in range(1, 14):    # A..M (13 столбцов)
            cell = ws.cell(row=row, column=col)
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border
            cell.fill = HEADER_FILL

    # ============================================================
    # ДАННЫЕ (со строки 3)
    # ============================================================
    MAX_PORT = 48
    position = 1
    current_row = 3

    # Сортируем по шкафу, панели и порту — чтобы строки шли по порядку
    df_sorted = df.sort_values(by=["Шкаф", "Панель", "Порт"]).reset_index(drop=True)

    # Группируем по (Шкаф, Панель)
    grouped = df_sorted.groupby(["Шкаф", "Панель"])

    for (cabinet, panel), panel_group in grouped:
        # Собираем карту занятых портов:
        #   {номер_порта: (имя, этаж, тип_блока, display_name)}
        taken = {}
        for _, r in panel_group.iterrows():
            taken[int(r["Порт"])] = (
                r["Имя"],
                int(r["Этаж"]),
                r["Тип"],
                r.get("_display_name", "") or r.get("Тип", ""),
            )

        # Проходим ВСЕ порты 1..48 (в том числе свободные)
        for port in range(1, MAX_PORT + 1):
            switch_port = switch_port_from_patch(port)

            if port in taken:
                # Порт занят оборудованием
                name, floor, dev_type, display = taken[port]
                type_str = display
                # Питание: PoE, если стоит галочка для этого типа блока
                poe_flag = poe_map.get(dev_type, False)
                power = "PoE" if poe_flag else "нет"
                # Маркировка по проекту — это имя, которое мы присвоили
                mark = name
                # Этаж — целое число
                floor_val = floor
                # IP и MAC — заглушка "-" (заполнит подрядчик)
                ip_val = "-"
                mac_val = "-"
            else:
                # Свободный (резервный) порт
                type_str = "резерв"
                mark = "-"
                floor_val = "-"
                power = "-"
                ip_val = "-"
                mac_val = "-"

            # Примечание — всегда пустое (заполнит подрядчик)
            note_val = ""

            # --- Записываем строку ---
            ws.cell(row=current_row, column=1, value=position)
            ws.cell(row=current_row, column=2, value=cabinet)
            ws.cell(row=current_row, column=3, value=panel)       # коммутатор = панель
            ws.cell(row=current_row, column=4, value=switch_port)
            ws.cell(row=current_row, column=5, value=panel)       # патч-панель
            ws.cell(row=current_row, column=6, value=port)
            ws.cell(row=current_row, column=7, value=type_str)
            ws.cell(row=current_row, column=8, value=mark)
            ws.cell(row=current_row, column=9, value=floor_val)
            ws.cell(row=current_row, column=10, value=power)
            ws.cell(row=current_row, column=11, value=ip_val)
            ws.cell(row=current_row, column=12, value=mac_val)
            ws.cell(row=current_row, column=13, value=note_val)

            # --- Стилизация строки ---
            for col in range(1, 14):
                cell = ws.cell(row=current_row, column=col)
                cell.border = thin_border
                # Числа и короткие поля — по центру
                if col in (1, 3, 4, 5, 6, 9, 10):
                    cell.alignment = center_align
                else:
                    cell.alignment = left_align

            position += 1
            current_row += 1

    # ============================================================
    # АВТОФИЛЬТР
    # ============================================================
    # Фильтр ставим на строку 2 (там конкретные заголовки столбцов),
    # последняя строка — последняя заполненная.
    last_data_row = current_row - 1
    ws.auto_filter.ref = f"A2:M{last_data_row}"

    # ============================================================
    # ШИРИНА КОЛОНОК
    # ============================================================
    widths = {
        'A': 6,    # №
        'B': 14,   # Шкаф №
        'C': 14,   # Коммутатор №
        'D': 18,   # Порт коммутатора №
        'E': 14,   # Патч-панель №
        'F': 18,   # Порт патч-панели №
        'G': 22,   # Тип оборудования
        'H': 24,   # Маркировка по проекту
        'I': 16,   # Этаж расположения
        'J': 10,   # Питание
        'K': 16,   # IP адрес
        'L': 16,   # MAC адрес
        'M': 22,   # Примечание
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    return ws
